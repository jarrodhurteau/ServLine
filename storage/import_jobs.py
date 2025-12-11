from __future__ import annotations
import csv
import json
import sqlite3
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

# Optional XLSX support (Phase 6 pt.3 — structured_xlsx ingestion)
try:
    import openpyxl  # type: ignore[import]
except Exception:  # pragma: no cover - defensive; library may not be installed
    openpyxl = None  # type: ignore[assignment]

ROOT = Path(__file__).resolve().parents[1]
DB_PATH = ROOT / "storage" / "servline.db"


# ---------------------------------------------------------------------------
# DB helpers
# ---------------------------------------------------------------------------

def db_connect() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def _get_import_jobs_columns() -> Set[str]:
    """
    Introspect the import_jobs table columns so we can insert flexibly.

    This lets us support new columns like:
      - source_type
      - restaurant_id
      - summary_json
      - payload_json

    without breaking older DBs that do not have them yet.
    """
    with db_connect() as conn:
        rows = conn.execute("PRAGMA table_info(import_jobs)").fetchall()
        return {str(r["name"]) for r in rows}


# ---------------------------------------------------------------------------
# Public read helpers
# ---------------------------------------------------------------------------

def get_import_job(job_id: int) -> Optional[Dict[str, Any]]:
    """
    Fetch a single import job entry by ID.

    Expected base schema (Day 20+):
        id INTEGER PRIMARY KEY
        filename TEXT
        source_path TEXT (uploads/<file>)
        status TEXT
        created_at TEXT
        updated_at TEXT
        source_type TEXT
        restaurant_id INTEGER
        summary_json TEXT
        payload_json TEXT

    Returns dict or None.
    """
    with db_connect() as conn:
        row = conn.execute(
            "SELECT * FROM import_jobs WHERE id = ?",
            (int(job_id),),
        ).fetchone()

        if not row:
            return None

        return {k: row[k] for k in row.keys()}


def list_import_jobs(limit: int = 200, offset: int = 0) -> List[Dict[str, Any]]:
    """
    Optional helper if you ever need it.
    """
    with db_connect() as conn:
        rows = conn.execute(
            "SELECT * FROM import_jobs ORDER BY id DESC LIMIT ? OFFSET ?",
            (limit, offset),
        ).fetchall()
        return [{k: r[k] for k in r.keys()} for r in rows]


# ---------------------------------------------------------------------------
# Structured ingestion: CSV/XLSX helpers
# ---------------------------------------------------------------------------

try:
    # One Brain contracts for structured ingestion (CSV / JSON / XLSX)
    from . import contracts as structured_contracts  # type: ignore[attr-defined]
except Exception:  # pragma: no cover - defensive; should exist in modern env
    structured_contracts = None  # type: ignore[assignment]


# Canonical field names for structured menu items.
CANONICAL_FIELDS = {
    "name",
    "description",
    "category",
    "subcategory",
    "price",
    "price_cents",
    "size",
    "sku",
}


HEADER_ALIASES: Dict[str, Set[str]] = {
    "name": {
        "name",
        "item",
        "itemname",
        "item_name",
        "itemtitle",
        "title",
        "menuitem",
    },
    "description": {
        "description",
        "desc",
        "details",
        "detail",
        "itemdescription",
    },
    "category": {
        "category",
        "cat",
        "section",
        "group",
        "menu_section",
        "menu_group",
    },
    "subcategory": {
        "subcategory",
        "subcat",
        "sub_category",
        "subsection",
        "sub_section",
    },
    "price": {
        "price",
        "cost",
        "amount",
        "baseprice",
        "base_price",
        "listprice",
    },
    "price_cents": {
        "pricecents",
        "price_cents",
    },
    "size": {
        "size",
        "portion",
        "variant",
        "serving",
    },
    "sku": {
        "sku",
        "code",
        "plu",
        "itemcode",
        "item_code",
    },
}


def _normalize_header(h: str) -> str:
    """
    Normalize a header to a simple token for alias matching.
    """
    return "".join(ch for ch in h.lower() if ch.isalnum())


def _detect_header_mapping(headers: List[str]) -> Dict[str, str]:
    """
    Detect which headers correspond to canonical structured fields.

    Returns mapping:

        { canonical_field: original_header }

    """
    mapping: Dict[str, str] = {}

    if not headers:
        return mapping

    normalized: Dict[str, str] = {h: _normalize_header(h) for h in headers}

    for canonical, aliases in HEADER_ALIASES.items():
        for header, norm in normalized.items():
            if norm in aliases:
                mapping[canonical] = header
                break

    return mapping


def _csv_row_to_structured_item(
    raw_row: Dict[str, Any],
    header_map: Dict[str, str],
) -> Dict[str, Any]:
    """
    Convert a raw row dict (CSV or XLSX or JSON-origin) into a canonical structured item dict.
    """
    item: Dict[str, Any] = {}

    for canonical in CANONICAL_FIELDS:
        source_header = header_map.get(canonical)
        if source_header is not None:
            item[canonical] = raw_row.get(source_header)

    # Collect leftover columns as meta
    meta: Dict[str, Any] = {}
    mapped_headers = set(header_map.values())
    for col_name, value in raw_row.items():
        if col_name not in mapped_headers:
            if value not in (None, "", " "):
                meta[col_name] = value

    if meta:
        item["meta"] = meta

    return item


def parse_structured_csv(
    csv_path: Path,
) -> Tuple[
    List[Dict[str, Any]],
    List[Dict[str, Any]],
    Dict[str, int],
    Dict[str, str],
    List[Dict[str, Any]],
]:
    """
    Parse a CSV file into structured menu items and run validation via One Brain contracts.

    Returns:
      clean_items:  list of normalized items ready for draft creation
      errors:       list of row-level validation error dicts
      summary:      summary dict (total_rows, valid_rows, error_rows)
      header_map:   mapping of canonical field to CSV header name
      raw_rows:     list of raw row dicts keyed by original headers (for mapping engine)
    """
    if structured_contracts is None:
        raise RuntimeError(
            "storage.contracts module is required for structured CSV parsing; "
            "make sure storage/contracts.py exists.",
        )

    csv_path = Path(csv_path)
    if not csv_path.exists():
        raise FileNotFoundError(csv_path)

    with csv_path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        headers = reader.fieldnames or []
        header_map = _detect_header_mapping(headers)

        raw_rows: List[Dict[str, Any]] = []
        structured_items: List[Dict[str, Any]] = []

        for row in reader:
            raw_row = dict(row)
            raw_rows.append(raw_row)
            structured_items.append(_csv_row_to_structured_item(raw_row, header_map))

    clean_items, errors, summary = structured_contracts.validate_structured_items(structured_items)
    return clean_items, errors, summary, header_map, raw_rows


def parse_structured_xlsx(
    xlsx_path: Path,
) -> Tuple[
    List[Dict[str, Any]],
    List[Dict[str, Any]],
    Dict[str, int],
    Dict[str, str],
    List[Dict[str, Any]],
]:
    """
    Parse an XLSX file (first sheet) into structured menu items and run validation
    via One Brain contracts.

    Behavior:
      - Uses the first worksheet in the workbook
      - Uses the first row as headers
      - Normalizes headers with the same logic as CSV
      - Reuses _csv_row_to_structured_item for row shaping

    Returns:
      clean_items:  list of normalized items ready for draft creation
      errors:       list of row-level validation error dicts
      summary:      summary dict (total_rows, valid_rows, error_rows)
      header_map:   mapping of canonical field to XLSX header name
      raw_rows:     list of raw row dicts keyed by original headers (for mapping engine)
    """
    if structured_contracts is None:
        raise RuntimeError(
            "storage.contracts module is required for structured XLSX parsing; "
            "make sure storage/contracts.py exists.",
        )

    if openpyxl is None:
        raise RuntimeError(
            "openpyxl is required for structured XLSX ingestion. "
            "Install it in your environment to enable XLSX imports.",
        )

    xlsx_path = Path(xlsx_path)
    if not xlsx_path.exists():
        raise FileNotFoundError(xlsx_path)

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)  # type: ignore[call-arg]
    sheets = wb.worksheets
    if not sheets:
        raise ValueError(f"XLSX file {xlsx_path} has no worksheets")

    ws = sheets[0]
    rows_iter = ws.iter_rows(values_only=True)

    # First row = headers
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise ValueError(f"XLSX file {xlsx_path} is empty or missing a header row")

    headers: List[str] = []
    for idx, cell in enumerate(header_row):
        if cell is None:
            headers.append(f"column_{idx + 1}")
        else:
            text = str(cell).strip()
            headers.append(text or f"column_{idx + 1}")

    header_map = _detect_header_mapping(headers)

    raw_rows: List[Dict[str, Any]] = []
    structured_items: List[Dict[str, Any]] = []

    for row_values in rows_iter:
        if row_values is None:
            continue

        row_dict: Dict[str, Any] = {}
        is_empty = True
        for idx, value in enumerate(row_values):
            if idx >= len(headers):
                break
            if value not in (None, "", " "):
                is_empty = False
            row_dict[headers[idx]] = value

        if is_empty:
            continue

        raw_rows.append(row_dict)
        structured_items.append(_csv_row_to_structured_item(row_dict, header_map))

    clean_items, errors, summary = structured_contracts.validate_structured_items(structured_items)
    return clean_items, errors, summary, header_map, raw_rows


# ---------------------------------------------------------------------------
# Structured import job creation (CSV/XLSX/JSON and generic)
# ---------------------------------------------------------------------------

def create_structured_import_job(
    source_type: str,
    *,
    filename: Optional[str] = None,
    source_path: Optional[str] = None,
    restaurant_id: Optional[int] = None,
    summary: Optional[Dict[str, Any]] = None,
    payload: Optional[Dict[str, Any]] = None,
    status: str = "parsed",
) -> int:
    """
    Create an import_jobs row for a structured import (CSV, JSON, XLSX).

    This function is resilient to schema differences by introspecting columns.

    Used by:
      - create_csv_import_job_from_file
      - create_xlsx_import_job_from_file
      - create_json_import_job_from_file

    Returns:
      job_id (int)
    """
    columns = _get_import_jobs_columns()
    insert_cols: List[str] = []
    values: List[Any] = []

    if "filename" in columns and filename is not None:
        insert_cols.append("filename")
        values.append(filename)

    if "source_path" in columns and source_path is not None:
        insert_cols.append("source_path")
        values.append(source_path)

    if "source_type" in columns:
        insert_cols.append("source_type")
        values.append(source_type)

    if "restaurant_id" in columns and restaurant_id is not None:
        insert_cols.append("restaurant_id")
        values.append(int(restaurant_id))

    if "status" in columns:
        insert_cols.append("status")
        values.append(status)

    if "summary_json" in columns and summary is not None:
        insert_cols.append("summary_json")
        values.append(json.dumps(summary))

    if "payload_json" in columns and payload is not None:
        insert_cols.append("payload_json")
        values.append(json.dumps(payload))

    if not insert_cols:
        raise RuntimeError("import_jobs table has no usable columns for INSERT")

    placeholders = ", ".join(["?"] * len(insert_cols))
    col_list = ", ".join(insert_cols)
    sql = f"INSERT INTO import_jobs ({col_list}) VALUES ({placeholders})"

    with db_connect() as conn:
        cur = conn.execute(sql, values)
        job_id = cur.lastrowid
        conn.commit()

    return int(job_id)


def create_csv_import_job_from_file(
    csv_path: Path,
    restaurant_id: Optional[int] = None,
) -> Dict[str, Any]:
    """
    High-level helper for Phase 6 pt.1.

    Steps:
      - Parse a CSV into structured items
      - Validate via One Brain contracts
      - Create an import_jobs row recording the results
      - Return details for draft creation and UI summary

    Returns dict:

      {
        "job_id": int,
        "items": [... clean items ...],
        "errors": [... error dicts ...],
        "summary": {... counts ...},
        "header_map": {... canonical to CSV header ...},
        "sample_rows": [... first N clean items ...],
        "job_summary": {... stored in summary_json (if column exists) ...},
      }

    """
    csv_path = Path(csv_path)
    clean_items, errors, summary, header_map, raw_rows = parse_structured_csv(csv_path)
    sample_rows = clean_items[:10]

    job_summary: Dict[str, Any] = {
        "ingest_mode": "structured_csv",
        "header_map": header_map,
        "summary": summary,
        "error_rows": errors,
        "sample_rows": sample_rows,
        "raw_rows": raw_rows,
    }

    # Store column_names for the mapping UI
    column_names = list(raw_rows[0].keys()) if raw_rows else []
    job_summary["column_names"] = column_names

    status = "parsed_with_errors" if errors else "parsed"

    job_id = create_structured_import_job(
        source_type="structured_csv",
        filename=csv_path.name,
        source_path=str(csv_path),
        restaurant_id=restaurant_id,
        summary=job_summary,
        payload=None,
        status=status,
    )

    return {
        "job_id": job_id,
        "items": clean_items,
        "errors": errors,
        "summary": summary,
        "header_map": header_map,
        "sample_rows": sample_rows,
        "job_summary": job_summary,
    }


def create_xlsx_import_job_from_file(
    xlsx_path: Path,
    restaurant_id: Optional[int] = None,
) -> Dict[str, Any]:
    """
    High-level helper for Phase 6 pt.3.

    Steps:
      - Parse an XLSX into structured items (first sheet, first row is headers)
      - Validate via One Brain contracts
      - Create an import_jobs row recording the results
      - Return details for draft creation and UI summary

    Returns dict:

      {
        "job_id": int,
        "items": [... clean items ...],
        "errors": [... error dicts ...],
        "summary": {... counts ...],
        "header_map": {... canonical to XLSX header ...},
        "sample_rows": [... first N clean items ...],
        "job_summary": {... stored in summary_json (if column exists) ...},
      }

    """
    xlsx_path = Path(xlsx_path)
    clean_items, errors, summary, header_map, raw_rows = parse_structured_xlsx(xlsx_path)
    sample_rows = clean_items[:10]

    job_summary: Dict[str, Any] = {
        "ingest_mode": "structured_xlsx",
        "header_map": header_map,
        "summary": summary,
        "error_rows": errors,
        "sample_rows": sample_rows,
        "raw_rows": raw_rows,
    }

    column_names = list(raw_rows[0].keys()) if raw_rows else []
    job_summary["column_names"] = column_names

    status = "parsed_with_errors" if errors else "parsed"

    job_id = create_structured_import_job(
        source_type="structured_xlsx",
        filename=xlsx_path.name,
        source_path=str(xlsx_path),
        restaurant_id=restaurant_id,
        summary=job_summary,
        payload=None,
        status=status,
    )

    return {
        "job_id": job_id,
        "items": clean_items,
        "errors": errors,
        "summary": summary,
        "header_map": header_map,
        "sample_rows": sample_rows,
        "job_summary": job_summary,
    }


def create_json_import_job_from_file(
    json_path: Path,
    restaurant_id: Optional[int] = None,
) -> Dict[str, Any]:
    """
    High-level helper for Phase 6 pt.7 (structured JSON).

    Steps:
      - Load a JSON file containing either:
          * {"items": [... structured or semi-structured items ...]}
          * [... structured or semi-structured items ...]
      - Validate via One Brain contracts
      - Create an import_jobs row recording the results
      - Return details for draft creation and UI summary

    For JSON, we also record:
      - raw_rows: original item dicts (for column mapping)
      - header_map: initial guess based on item keys (may be empty)

    Returns dict:

      {
        "job_id": int,
        "items": [... clean items ...],
        "errors": [... error dicts ...],
        "summary": {... counts ...},
        "header_map": {...},
        "sample_rows": [... first N clean items ...],
        "job_summary": {... stored in summary_json (if column exists) ...},
      }

    """
    if structured_contracts is None:
        raise RuntimeError(
            "storage.contracts module is required for structured JSON parsing; "
            "make sure storage/contracts.py exists.",
        )

    json_path = Path(json_path)
    if not json_path.exists():
        raise FileNotFoundError(json_path)

    try:
        with json_path.open("r", encoding="utf-8") as f:
            payload = json.load(f)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Invalid JSON file {json_path.name}: {exc}") from exc

    clean_items: List[Dict[str, Any]] = []
    errors: List[Dict[str, Any]] = []
    summary: Dict[str, int] = {}

    raw_rows: List[Dict[str, Any]] = []
    if isinstance(payload, dict):
        items = payload.get("items")
        if isinstance(items, list):
            raw_rows = [dict(row) for row in items if isinstance(row, dict)]
    elif isinstance(payload, list):
        raw_rows = [dict(row) for row in payload if isinstance(row, dict)]

    header_map: Dict[str, str] = {}
    if raw_rows:
        first_row = raw_rows[0]
        headers = list(first_row.keys())
        header_map = _detect_header_mapping(headers)

    if isinstance(payload, dict):
        ok, msg, clean_items, item_errors, base_summary = structured_contracts.validate_structured_menu_payload(payload)
        errors = item_errors
        summary = base_summary or {
            "total_rows": len(clean_items) + len(errors),
            "valid_rows": len(clean_items),
            "error_rows": len(errors),
        }

        if not ok and summary.get("total_rows", 0) == 0:
            raise ValueError(f"Invalid structured JSON payload: {msg}")
    elif isinstance(payload, list):
        clean_items, item_errors, base_summary = structured_contracts.validate_structured_items(payload)
        errors = item_errors
        summary = base_summary or {
            "total_rows": len(payload),
            "valid_rows": len(clean_items),
            "error_rows": len(errors),
        }
    else:
        raise ValueError(
            "Structured JSON payload must be either an object with an 'items' array "
            "or a top-level list of items.",
        )

    sample_rows = clean_items[:10]

    job_summary: Dict[str, Any] = {
        "ingest_mode": "structured_json",
        "header_map": header_map,
        "summary": summary,
        "error_rows": errors,
        "sample_rows": sample_rows,
        "raw_rows": raw_rows,
    }

    column_names = list(raw_rows[0].keys()) if raw_rows else []
    job_summary["column_names"] = column_names

    status = "parsed_with_errors" if errors else "parsed"

    job_id = create_structured_import_job(
        source_type="structured_json",
        filename=json_path.name,
        source_path=str(json_path),
        restaurant_id=restaurant_id,
        summary=job_summary,
        payload=payload if isinstance(payload, (dict, list)) else None,
        status=status,
    )

    return {
        "job_id": job_id,
        "items": clean_items,
        "errors": errors,
        "summary": summary,
        "header_map": header_map,
        "sample_rows": sample_rows,
        "job_summary": job_summary,
    }


# ---------------------------------------------------------------------------
# Column mapping and rebuild helpers (Phase 6 pt.9)
# ---------------------------------------------------------------------------

def get_structured_job_summary(job_id: int) -> Tuple[Dict[str, Any], Dict[str, Any]]:
    """
    Convenience helper for mapping UI and engine.

    Returns:
      job:     full import_jobs row as a dict
      summary: parsed summary_json dict (or empty dict)
    """
    job = get_import_job(job_id)
    if not job:
        raise ValueError(f"import job {job_id} not found")

    raw_summary = job.get("summary_json")
    summary: Dict[str, Any]
    if raw_summary:
        try:
            summary = json.loads(raw_summary)
        except Exception:
            summary = {}
    else:
        summary = {}

    return job, summary


def update_import_job_header_map(job_id: int, header_map: Dict[str, str]) -> None:
    """
    Persist an updated header_map into import_jobs.summary_json.

    This does not recompute items; it only writes the metadata.
    """
    columns = _get_import_jobs_columns()
    if "summary_json" not in columns:
        return

    job, summary = get_structured_job_summary(job_id)
    if not isinstance(summary, dict):
        summary = {}

    summary["header_map"] = header_map

    with db_connect() as conn:
        conn.execute(
            "UPDATE import_jobs SET summary_json = ? WHERE id = ?",
            (json.dumps(summary), int(job_id)),
        )
        conn.commit()


def rebuild_structured_items_from_header_map(
    job_id: int,
    header_map: Optional[Dict[str, str]] = None,
) -> Tuple[
    List[Dict[str, Any]],
    List[Dict[str, Any]],
    Dict[str, int],
    Dict[str, str],
    List[Dict[str, Any]],
]:
    """
    Rebuild structured items for a job using an updated header_map.

    This powers the Column Mapping engine.

    Steps:
      - Load raw_rows from summary_json (original CSV, XLSX, or JSON rows)
      - Apply the supplied header_map (or existing one if not provided)
      - Run validate_structured_items on the mapped rows via One Brain contracts
      - Update summary_json with the new mapping and preview data

    Returns:
      clean_items:  new normalized items ready for replacing draft items
      errors:       list of validation error dicts
      summary:      summary dict
      header_map:   the header_map that was used (final)
      sample_rows:  first N clean items
    """
    if structured_contracts is None:
        raise RuntimeError(
            "storage.contracts module is required for structured mapping; "
            "make sure storage/contracts.py exists.",
        )

    columns = _get_import_jobs_columns()
    job, summary = get_structured_job_summary(job_id)

    raw_rows = summary.get("raw_rows") or []
    if not isinstance(raw_rows, list) or not raw_rows:
        raise ValueError(
            "This import job does not have raw_rows recorded. "
            "Column mapping based rebuild is not available.",
        )

    if header_map is None:
        header_map = summary.get("header_map") or {}
    if not isinstance(header_map, dict):
        header_map = {}

    structured_items: List[Dict[str, Any]] = []
    for raw_row in raw_rows:
        if not isinstance(raw_row, dict):
            continue
        structured_items.append(_csv_row_to_structured_item(raw_row, header_map))

    clean_items, errors, base_summary = structured_contracts.validate_structured_items(structured_items)
    summary_counts = base_summary or {
        "total_rows": len(structured_items),
        "valid_rows": len(clean_items),
        "error_rows": len(errors),
    }
    sample_rows = clean_items[:10]

    summary["header_map"] = header_map
    summary["summary"] = summary_counts
    summary["error_rows"] = errors
    summary["sample_rows"] = sample_rows

    if "summary_json" in columns:
        with db_connect() as conn:
            conn.execute(
                "UPDATE import_jobs SET summary_json = ? WHERE id = ?",
                (json.dumps(summary), int(job_id)),
            )
            conn.commit()

    return clean_items, errors, summary_counts, header_map, sample_rows
